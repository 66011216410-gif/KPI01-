import io
import re
import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

PAGE_TITLE='ฐานข้อมูลจำนวนผลงานระดับบัณฑิตศึกษาที่ตีพิมพ์ในระดับนานาชาติ Scopus Q1, Q2'
st.set_page_config(page_title=PAGE_TITLE, page_icon='📊', layout='wide')
LEVEL=['ระดับ','ระดับการศึกษา','degree']; BASE=['ฐานข้อมูล','ฐานข้อมูล ']; GROUP=['กลุ่มสาขา']; FACULTY=['คณะ','faculty']; PROGRAM_TOTAL=['สาขารวม','สาขา รวม','program total']; PROGRAM=['สาขา','สาขาวิชา','หลักสูตร','program','major']; SYSTEM=['ระบบ']; STUDENT=['รหัสนิสิต','รหัสนักศึกษา','รหัส','student id']

def clean(v): return '' if pd.isna(v) else str(v).strip()
def norm_col(v): return re.sub(r'\s+','',clean(v)).lower()
def find_col(df,aliases):
    cols={norm_col(c):c for c in df.columns}
    for a in aliases:
        if norm_col(a) in cols:return cols[norm_col(a)]
    for c in df.columns:
        if any(norm_col(a) in norm_col(c) for a in aliases):return c
    return None
def norm_level(v):
    s=clean(v).lower().replace(' ','')
    if any(x in s for x in ['ปริญญาเอก','ป.เอก','phd','ph.d','doctoral']):return 'ปริญญาเอก'
    if any(x in s for x in ['ปริญญาโท','ป.โท','master']):return 'ปริญญาโท'
    return ''
def level_from_sheet(name):
    s=clean(name).lower().replace(' ','')
    if any(x in s for x in ['เอก','ปริญญาเอก','phd','doctoral']):return 'ปริญญาเอก'
    if any(x in s for x in ['โท','ปริญญาโท','master']):return 'ปริญญาโท'
    return ''
def norm_q(v):
    s=clean(v).upper().replace('SCOPUS(Q','SCOPUS (Q')
    if 'SCOPUS (Q1)' in s:return 'Q1'
    if 'SCOPUS (Q2)' in s:return 'Q2'
    return ''
def norm_system(v):
    s=re.sub(r'\s+','',clean(v)).lower()
    if s.startswith('ระบบ'):s=s[4:]
    if any(x in s for x in ['นอกเวลาราชการ','นอกเวลา','parttime','part-time']):return 'นอกเวลา'
    if any(x in s for x in ['ในเวลาราชการ','ในเวลา','fulltime','full-time']):return 'ในเวลา'
    return ''
def prepare_sheet(df,sheet_name):
    out=df.copy();out.columns=[clean(c) for c in out.columns]
    level_col=find_col(out,LEVEL);base_col=find_col(out,BASE);group_col=find_col(out,GROUP);faculty_col=find_col(out,FACULTY);program_total_col=find_col(out,PROGRAM_TOTAL);program_col=find_col(out,PROGRAM);system_col=find_col(out,SYSTEM);student_col=find_col(out,STUDENT)
    required=[]
    if not base_col:required.append('ฐานข้อมูล')
    if not group_col:required.append('กลุ่มสาขา')
    if not system_col:required.append('ระบบ')
    if required:raise ValueError(f"Sheet '{sheet_name}' ไม่พบคอลัมน์: {', '.join(required)}")
    forced=level_from_sheet(sheet_name)
    level=pd.Series(forced,index=out.index) if forced else (out[level_col].map(norm_level) if level_col else pd.Series('',index=out.index))
    group=out[group_col].map(clean).replace('',pd.NA).ffill().fillna('')
    if faculty_col:
        raw=out[faculty_col].map(clean).replace('',pd.NA);vals=[];lg=None;lf=''
        for g,f in zip(group,raw):
            g=clean(g)
            if g!=lg:lg,lf=g,''
            if pd.notna(f) and clean(f):lf=clean(f)
            vals.append(lf)
        faculty=pd.Series(vals,index=out.index)
    else:faculty=pd.Series('',index=out.index)
    source=out[program_total_col] if program_total_col else (out[program_col] if program_col else None)
    if source is not None:
        raw=source.map(clean).replace('',pd.NA);vals=[];lg=None;lf=None;lp=''
        for g,f,p in zip(group,faculty,raw):
            g,f=clean(g),clean(f)
            if g!=lg or f!=lf:lg,lf,lp=g,f,''
            if pd.notna(p) and clean(p):lp=clean(p)
            vals.append(lp)
        program=pd.Series(vals,index=out.index)
    else:program=pd.Series('',index=out.index)
    faculty_group={};program_group={}
    for g,f,p in zip(group,faculty,program):
        g,f,p=clean(g),clean(f),clean(p)
        if g and f and f not in faculty_group:faculty_group[f]=g
        if g and p and p not in program_group:program_group[p]=g
    canonical=[]
    for g,f,p in zip(group,faculty,program):
        g,f,p=clean(g),clean(f),clean(p)
        if f and f in faculty_group:g=faculty_group[f]
        elif p and p in program_group:g=program_group[p]
        canonical.append(g)
    out['__ระดับ']=level;out['__กลุ่ม']=canonical;out['__คณะ']=faculty.map(clean);out['__สาขา']=program.map(clean);out['__Q']=out[base_col].map(norm_q);out['__ระบบ']=out[system_col].map(norm_system);out['__รหัส']=out[student_col].map(clean) if student_col else '';out['__ลำดับ']=range(len(out))
    return out[out['__ระดับ'].isin(['ปริญญาโท','ปริญญาเอก'])].copy()
def read_excel(uploaded):
    uploaded.seek(0);sheets=pd.read_excel(uploaded,sheet_name=None,header=1);frames=[];used=[];raw_sheets={};errors=[]
    for name,df in sheets.items():
        if df is None or df.empty:continue
        try:
            p=prepare_sheet(df,name)
            if not p.empty:frames.append(p);used.append(name);raw_sheets[name]=df.copy()
        except ValueError as e:errors.append(str(e))
    if not frames:raise ValueError('ไม่พบ Sheet ข้อมูลปริญญาโท/ปริญญาเอกที่ใช้งานได้'+(f": {'; '.join(errors[:3])}" if errors else ''))
    return used,pd.concat(frames,ignore_index=True),raw_sheets
def count_system(df,kind):
    x=df[df['__ระบบ']==kind]
    if x.empty:return 0
    ids=x['__รหัส'].map(clean)
    return int(ids[ids!=''].nunique()+int((ids=='').sum()))
def metrics(df):
    inn=count_system(df,'ในเวลา');out=count_system(df,'นอกเวลา');a=inn+out;pub=int(df['__Q'].isin(['Q1','Q2']).sum());return inn,out,a,pub,round(pub*100/a,2) if a else 0
ORDER_TREE=[('กลุ่มมนุษยศาสตร์และสังคมศาสตร์',[('การเมืองการปกครอง',['รัฐศาสตร์']),('การท่องเที่ยวและการโรงแรม',['การจัดการการท่องเที่ยวและการโรงแรม']),('การบัญชีและการจัดการ',['การจัดการสมัยใหม่','การจัดการสมาร์ตซิตี้และนวัตกรรมดิจิทัล','การบัญชี','บริหารธุรกิจและนวัตกรรมดิจิทัล']),('ดุริยางคศิลป์',['ดุริยางคศิลป์']),('มนุษยศาสตร์และสังคมศาสตร์',['การสอนภาษาอังกฤษ','ภาษาไทย','ศาสนาและภูมิปัญญาเพื่อการพัฒนา']),('ศิลปกรรมศาสตร์และวัฒนธรรมศาสตร์',['การวิจัยและสร้างสรรค์ศิลปกรรมศาสตร์','วัฒนธรรมศาสตร์']),('ศึกษาศาสตร์',['เทคโนโลยีและสื่อสารการศึกษา','การบริหารและพัฒนาการศึกษา','วิจัยและประเมินผลการศึกษา','วิทยาศาสตร์การออกกำลังกายและการกีฬา','หลักสูตรและการสอน'])]),('กลุ่มวิทยาศาสตร์และเทคโนโลยี',[('เทคโนโลยี',['เกษตรศาสตร์','เทคโนโลยีการอาหาร']),('วิจัยวลัยรุกขเวช',['ความหลากหลายทางชีวภาพ']),('วิทยาการสารสนเทศ',['เทคโนโลยีสารสนเทศ','วิทยาการคอมพิวเตอร์','สื่อนฤมิต']),('วิทยาศาสตร์',['บรรพชีวินวิทยา','ฟิสิกส์']),('วิศวกรรมศาสตร์',['วิศวกรรมเครื่องกล','วิศวกรรมโยธา','วิศวกรรมไฟฟ้าและคอมพิวเตอร์']),('สิ่งแวดล้อมและทรัพยากรศาสตร์',['การจัดการสิ่งแวดล้อมอย่างยั่งยืน'])]),('กลุ่มวิทยาศาสตร์สุขภาพ',[('เภสัชศาสตร์',['เภสัชศาสตร์']),('แพทยศาสตร์',['วิทยาศาสตร์สุขภาพ']),('สาธารณสุขศาสตร์',['เทคโนโลยีทางสุขภาพและความปลอดภัย','สาธารณสุขศาสตรดุษฎีบัณฑิต'])])]
def hierarchy(data):
    cols=['ลำดับ','กลุ่ม/คณะ/สาขา','ระบบในเวลาราชการ','ระบบนอกเวลาราชการ','จำนวนผู้สำเร็จการศึกษา (A)','รวมจำนวนผลงานตีพิมพ์ระดับนานาชาติ Q1-Q2','ร้อยละ'];rows=[];seen_g=set();seen_f=set();seen_p=set();n=1
    def add_group(g):
        nonlocal n
        if not g or g in seen_g:return
        m=data['__กลุ่ม'].map(clean).eq(g)
        if m.any():seen_g.add(g);rows.append([n,g,*metrics(data[m])]);n+=1
    def add_faculty(g,f):
        nonlocal n
        key=(g,f);m=data['__กลุ่ม'].map(clean).eq(g)&data['__คณะ'].map(clean).eq(f)
        if f and key not in seen_f and m.any():seen_f.add(key);rows.append([n,'    '+f,*metrics(data[m])]);n+=1
    def add_program(g,f,p):
        nonlocal n
        key=(g,f,p);m=data['__กลุ่ม'].map(clean).eq(g)&data['__คณะ'].map(clean).eq(f)&data['__สาขา'].map(clean).eq(p)
        if p and key not in seen_p and m.any():seen_p.add(key);rows.append([n,'        '+p,*metrics(data[m])]);n+=1
    def programs_for_faculty(g,f,preferred):
        mask=data['__กลุ่ม'].map(clean).eq(g)&data['__คณะ'].map(clean).eq(f)&data['__สาขา'].map(clean).ne('');actual=[];seen=set()
        for p in data.loc[mask,'__สาขา'].map(clean):
            if p and p not in seen:seen.add(p);actual.append(p)
        ordered=[p for p in preferred if p in seen];ordered.extend(p for p in actual if p not in ordered);return ordered
    for g,fs in ORDER_TREE:
        add_group(g)
        for f,ps in fs:
            add_faculty(g,f)
            for p in programs_for_faculty(g,f,ps):add_program(g,f,p)
    for _,r in data.sort_values('__ลำดับ').iterrows():
        g,f=clean(r['__กลุ่ม']),clean(r['__คณะ']);add_group(g)
        if g and f:
            add_faculty(g,f)
            for p in programs_for_faculty(g,f,[]):add_program(g,f,p)
    rows.append(['','รวมทั้งหมด',*metrics(data)]);return pd.DataFrame(rows,columns=cols)
def q_summary(data):
    def vals(level):
        x=data[data['__ระดับ']==level];q1=int((x['__Q']=='Q1').sum());q2=int((x['__Q']=='Q2').sum());return q1,q2,q1+q2
    m,d=vals('ปริญญาโท'),vals('ปริญญาเอก');return pd.DataFrame([['ระดับปริญญาโท',*m],['ระดับปริญญาเอก',*d],['รวม',m[0]+d[0],m[1]+d[1],m[2]+d[2]]],columns=['','Scopus Q1','Scopus Q2','รวม'])
def safe_sheet_name(name,used):
    base='ข้อมูล_'+clean(name)[:25];base=re.sub(r'[\\/*?:\[\]]','_',base) or 'ข้อมูล';candidate=base;i=2
    while candidate in used:candidate=f'{base[:28-len(str(i))]}_{i}';i+=1
    used.add(candidate);return candidate
def make_excel(data,raw_sheets):
    bio=io.BytesIO()
    with pd.ExcelWriter(bio,engine='openpyxl') as w:
        for level in ['ปริญญาโท','ปริญญาเอก']:hierarchy(data[data['__ระดับ']==level]).to_excel(w,sheet_name=level,index=False,startrow=3)
        q_summary(data).to_excel(w,sheet_name='สรุป Scopus Q1-Q2',index=False)
        used={'ปริญญาโท','ปริญญาเอก','สรุป Scopus Q1-Q2'}
        for name,raw in raw_sheets.items():raw.to_excel(w,sheet_name=safe_sheet_name(name,used),index=False)
    bio.seek(0);wb=load_workbook(bio);thin=Side(style='thin',color='000000');table_border=Border(left=thin,right=thin,top=thin,bottom=thin)
    for ws in wb.worksheets:
        if ws.title in ['ปริญญาโท','ปริญญาเอก']:
            # หัวตาราง 3 ชั้น ให้ตรงกับรูปแบบที่ต้องการ
            ws.insert_rows(1,3)
            # ลบแถวว่าง 4-6 แล้วเลื่อนข้อมูลขึ้นมาให้เริ่มที่แถว 4
            ws.delete_rows(4,3)
            ws.merge_cells('A1:G1')
            ws['A1']=PAGE_TITLE+' ระดับ'+ws.title
            ws['A1'].font=Font(bold=True,color='FFFFFF',size=16)
            ws['A1'].fill=PatternFill('solid',fgColor='2F5D20')
            ws['A1'].alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            ws.row_dimensions[1].height=28

            # แถวหัวตารางหลัก
            ws.merge_cells('A2:A3')
            ws.merge_cells('B2:B3')
            ws.merge_cells('C2:D2')
            ws.merge_cells('E2:E3')
            ws.merge_cells('F2:G2')

            ws['A2']='ลำดับ'
            ws['B2']='กลุ่ม/คณะ/สาขา'
            ws['C2']='ระบบ'
            ws['E2']='จำนวนผู้สำเร็จ\nการศึกษา (A)'
            ws['F2']='1.2.4 จำนวนผลงานระดับบัณฑิตศึกษาที่\nสามารถตีพิมพ์ในฐานข้อมูลนานาชาติ SCOPUS\nQ1-Q2'

            # แถวหัวตารางย่อย
            ws['C3']='ระบบในเวลา\nราชการ'
            ws['D3']='ระบบนอกเวลา\nราชการ'
            ws['F3']='รวมจำนวนผลงานตีพิมพ์\nระดับนานาชาติ Q1-Q2'
            ws['G3']='ร้อยละ'

            # เติมสีและจัดรูปแบบหัวตาราง
            for row in range(2,4):
                for c in range(1,8):
                    ws.cell(row,c).fill=PatternFill('solid',fgColor='2F5D20' if row == 2 else '568B3B')
                    ws.cell(row,c).font=Font(bold=True,color='FFFFFF')
                    ws.cell(row,c).alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
                    ws.cell(row,c).border=table_border

            ws.row_dimensions[2].height=54
            ws.row_dimensions[3].height=54
            ws.freeze_panes='A4'
            # ลบคอลัมน์ที่ 4-7 (D:G) ออกจากตาราง Excel
            ws.delete_cols(4,4)
            # ลบแถวว่างทั้งหมดในตาราง เพื่อให้ข้อมูลต่อเนื่องไม่มีช่องว่าง
            for r in range(ws.max_row,3,-1):
                if all(ws.cell(r,c).value in (None,'') for c in range(1,8)):
                    ws.delete_rows(r,1)
            ws.column_dimensions['A'].width=9
            ws.column_dimensions['B'].width=48
            for col in range(3,8):ws.column_dimensions[chr(64+col)].width=22
            for row in range(4,ws.max_row+1):
                label=str(ws.cell(row,2).value or '')
                if label.startswith('        '):ws.cell(row,2).value=label.strip();indent=2;fill='FFFFFF'
                elif label.startswith('    '):ws.cell(row,2).value=label.strip();indent=1;fill='E2F0D9'
                elif label=='รวมทั้งหมด':fill='FCE4D6';indent=0
                else:fill='FCE4D6';indent=0
                for c in range(1,8):ws.cell(row,c).fill=PatternFill('solid',fgColor=fill);ws.cell(row,c).border=table_border
                ws.cell(row,2).alignment=Alignment(horizontal='left',vertical='center',indent=indent)
        elif ws.title=='สรุป Scopus Q1-Q2':
            ws.freeze_panes='A2'
            for row in ws.iter_rows():
                for cell in row:
                    cell.border=table_border
    out=io.BytesIO();wb.save(out);out.seek(0);return out

st.title(PAGE_TITLE)
st.caption('ระบบจัดทำสถิติผลงานระดับบัณฑิตศึกษา Scopus Q1-Q2 แยกปริญญาโทและปริญญาเอก')
file=st.file_uploader('อัปโหลดไฟล์ Excel',type=['xlsx','xls'])
if file:
    try:
        used,data,raw=read_excel(file);st.success('อ่านข้อมูลสำเร็จ: '+', '.join(used))
        for level in ['ปริญญาโท','ปริญญาเอก']:
            st.subheader(level);st.dataframe(hierarchy(data[data['__ระดับ']==level]),use_container_width=True,hide_index=True)
        st.subheader('สรุป Scopus Q1-Q2');st.dataframe(q_summary(data),use_container_width=True,hide_index=True)
        st.download_button('ดาวน์โหลด Excel สถิติผลงานตีพิมพ์',data=make_excel(data,raw).getvalue(),file_name='สถิติผลงานตีพิมพ์.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e:st.error(f'ไม่สามารถอ่านไฟล์ได้: {e}')
else:st.info('อัปโหลดไฟล์ Excel เพื่อสร้างตารางสถิติ')
