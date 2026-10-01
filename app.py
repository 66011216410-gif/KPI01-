import io
import re
import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

st.set_page_config(page_title='KPI01', page_icon='📊', layout='wide')
LEVEL=['ระดับ','ระดับการศึกษา','degree']; Q_MAIN=['ฐานข้อมูล','ฐานข้อมูล ']; GROUP=['กลุ่มสาขา']; FACULTY=['คณะ','faculty']; PROGRAM_TOTAL=['สาขารวม','สาขา รวม','program total']; PROGRAM=['สาขา','สาขาวิชา','หลักสูตร','program','major']; SYSTEM=['ระบบ']; STUDENT=['รหัสนิสิต','รหัสนักศึกษา','รหัส','student id']

def clean(v): return '' if pd.isna(v) else str(v).strip()
def norm_col(s): return re.sub(r'\s+','',clean(s)).lower()
def find_col(df,aliases):
    cols={norm_col(c):c for c in df.columns}
    for a in aliases:
        if norm_col(a) in cols:return cols[norm_col(a)]
    for c in df.columns:
        if any(norm_col(a) in norm_col(c) for a in aliases):return c
    return None

def norm_level(v):
    s=clean(v).lower().replace(' ','')
    if any(x in s for x in ['ปริญญาเอก','ป.เอก','phd','ph.d','doctoral']): return 'ปริญญาเอก'
    if any(x in s for x in ['ปริญญาโท','ป.โท','master']): return 'ปริญญาโท'
    return 'อื่น ๆ'

def norm_q(v):
    s=clean(v).upper().replace('SCOPUS(Q','SCOPUS (Q')
    if 'SCOPUS (Q1)' in s:return 'Q1'
    if 'SCOPUS (Q2)' in s:return 'Q2'
    return 'อื่น ๆ'

def norm_system(v):
    s=re.sub(r'\s+','',clean(v)).lower()
    if s.startswith('ระบบ'):s=s[4:]
    if any(x in s for x in ['นอกเวลาราชการ','นอกเวลา','นอกเวล','parttime','part-time']):return 'นอกเวลา'
    if any(x in s for x in ['ในเวลาราชการ','ในเวลา','ในเวล','fulltime','full-time']):return 'ในเวลา'
    return ''

def read_excel(uploaded):
    uploaded.seek(0); sheets=pd.read_excel(uploaded,sheet_name=None,header=1); candidates=[]
    for name,df in sheets.items():
        if df is None or df.empty:continue
        df=df.copy();df.columns=[clean(c) for c in df.columns]
        if not (find_col(df,LEVEL) and find_col(df,Q_MAIN) and find_col(df,GROUP) and find_col(df,SYSTEM)):continue
        score=100+(10 if find_col(df,FACULTY) else 0)+(10 if (find_col(df,PROGRAM_TOTAL) or find_col(df,PROGRAM)) else 0);candidates.append((score,name,df))
    if not candidates:raise ValueError("ไม่พบ Sheet ที่มีคอลัมน์ ระดับ, ฐานข้อมูล, ระบบ และ กลุ่มสาขา ครบ")
    candidates.sort(key=lambda x:x[0],reverse=True);return candidates[0][1],candidates[0][2]

def prepare(df):
    out=df.copy(); level_col=find_col(out,LEVEL);base_col=find_col(out,Q_MAIN);group_col=find_col(out,GROUP);faculty_col=find_col(out,FACULTY);program_total_col=find_col(out,PROGRAM_TOTAL);program_col=find_col(out,PROGRAM);system_col=find_col(out,SYSTEM);student_col=find_col(out,STUDENT)
    if not level_col:raise ValueError("ไม่พบคอลัมน์ 'ระดับ' ในแถวที่ 2")
    if not base_col:raise ValueError("ไม่พบคอลัมน์ 'ฐานข้อมูล' ในแถวที่ 2")
    if not system_col:raise ValueError("ไม่พบคอลัมน์ 'ระบบ' ในแถวที่ 2")
    if not group_col:raise ValueError("ไม่พบคอลัมน์ 'กลุ่มสาขา' ในแถวที่ 2")
    raw_group=out[group_col].map(clean).replace('',pd.NA).ffill().fillna('')
    if faculty_col:
        fv=out[faculty_col].map(clean).replace('',pd.NA); vals=[]; cg=None; cf=''
        for g,f in zip(raw_group,fv):
            g=clean(g)
            if g!=cg: cg=g;cf=''
            if pd.notna(f) and clean(f)!='': cf=clean(f)
            vals.append(cf)
        raw_faculty=pd.Series(vals,index=out.index)
    else: raw_faculty=pd.Series('',index=out.index)
    ps=out[program_total_col] if program_total_col else (out[program_col] if program_col else None)
    if ps is not None:
        pv=ps.map(clean).replace('',pd.NA); vals=[]; cg=None; cf=None; cp=''
        for g,f,p in zip(raw_group,raw_faculty,pv):
            g=clean(g);f=clean(f)
            if g!=cg or f!=cf: cg=g;cf=f;cp=''
            if pd.notna(p) and clean(p)!='': cp=clean(p)
            vals.append(cp)
        raw_program=pd.Series(vals,index=out.index)
    else: raw_program=pd.Series('',index=out.index)
    # คณะหนึ่งคณะและสาขาหนึ่งสาขาจะถูกผูกกับกลุ่มเดียว โดยใช้กลุ่มที่พบครั้งแรกจากกลุ่มสาขา
    faculty_group={}; program_group={}
    for g,f,p in zip(raw_group,raw_faculty,raw_program):
        g=clean(g);f=clean(f);p=clean(p)
        if g and f and f not in faculty_group: faculty_group[f]=g
        if g and p and p not in program_group: program_group[p]=g
    canonical_group=[]
    for g,f,p in zip(raw_group,raw_faculty,raw_program):
        g=clean(g);f=clean(f);p=clean(p)
        if f and f in faculty_group:g=faculty_group[f]
        elif p and p in program_group:g=program_group[p]
        canonical_group.append(g)
    out['__กลุ่ม']=canonical_group;out['__คณะ']=raw_faculty.map(clean);out['__สาขา']=raw_program.map(clean)
    out['__ระดับ']=out[level_col].map(norm_level);out['__Q']=out[base_col].map(norm_q);out['__ระบบ']=out[system_col].map(norm_system);out['__รหัส']=out[student_col].map(clean) if student_col else '';out['__ลำดับ']=range(len(out))
    return out[out['__ระดับ'].isin(['ปริญญาโท','ปริญญาเอก'])].copy()

def count_system(df,kind):
    x=df[df['__ระบบ']==kind]
    if x.empty:return 0
    ids=x['__รหัส'].map(clean);return int(ids[ids!=''].nunique()+int((ids=='').sum()))

def metrics(df):
    inn=count_system(df,'ในเวลา');outn=count_system(df,'นอกเวลา');a=inn+outn;pub=int(df['__Q'].isin(['Q1','Q2']).sum());return inn,outn,a,pub,round(pub*100/a,2) if a else 0

ORDER_TREE=[
 ('กลุ่มมนุษยศาสตร์และสังคมศาสตร์',[('การเมืองการปกครอง',['รัฐศาสตร์']),('การท่องเที่ยวและการโรงแรม',['การจัดการการท่องเที่ยวและการโรงแรม']),('การบัญชีและการจัดการ',['การจัดการสมัยใหม่','การจัดการสมาร์ตซิตี้และนวัตกรรมดิจิทัล','การบัญชี','บริหารธุรกิจและนวัตกรรมดิจิทัล']),('ดุริยางคศิลป์',['ดุริยางคศิลป์']),('มนุษยศาสตร์และสังคมศาสตร์',['การสอนภาษาอังกฤษ','ภาษาไทย','ศาสนาและภูมิปัญญาเพื่อการพัฒนา']),('ศิลปกรรมศาสตร์และวัฒนธรรมศาสตร์',['การวิจัยและสร้างสรรค์ศิลปกรรมศาสตร์','วัฒนธรรมศาสตร์']),('ศึกษาศาสตร์',['เทคโนโลยีและสื่อสารการศึกษา','การบริหารและพัฒนาการศึกษา','วิจัยและประเมินผลการศึกษา','วิทยาศาสตร์การออกกำลังกายและการกีฬา','หลักสูตรและการสอน'])]),
 ('กลุ่มวิทยาศาสตร์และเทคโนโลยี',[('เทคโนโลยี',['เกษตรศาสตร์','เทคโนโลยีการอาหาร']),('วิจัยวลัยรุกขเวช',['ความหลากหลายทางชีวภาพ']),('วิทยาการสารสนเทศ',['เทคโนโลยีสารสนเทศ','วิทยาการคอมพิวเตอร์','สื่อนฤมิต']),('วิทยาศาสตร์',['บรรพชีวินวิทยา','ฟิสิกส์']),('วิศวกรรมศาสตร์',['วิศวกรรมเครื่องกล','วิศวกรรมโยธา','วิศวกรรมไฟฟ้าและคอมพิวเตอร์']),('สิ่งแวดล้อมและทรัพยากรศาสตร์',['การจัดการสิ่งแวดล้อมอย่างยั่งยืน'])]),
 ('กลุ่มวิทยาศาสตร์สุขภาพ',[('เภสัชศาสตร์',['เภสัชศาสตร์']),('แพทยศาสตร์',['วิทยาศาสตร์สุขภาพ']),('สาธารณสุขศาสตร์',['เทคโนโลยีทางสุขภาพและความปลอดภัย','สาธารณสุขศาสตรดุษฎีบัณฑิต'])])]

def hierarchy(data):
    cols=['ลำดับ','กลุ่ม/คณะ/สาขา','ระบบในเวลาราชการ','ระบบนอกเวลาราชการ','จำนวนผู้สำเร็จการศึกษา (A)','รวมจำนวนผลงานตีพิมพ์ระดับนานาชาติ Q1-Q2','ร้อยละ'];rows=[];seen_g=set();seen_f=set();seen_p=set();n=1
    def add_group(g):
        nonlocal n
        if not g or g in seen_g:return
        m=data['__กลุ่ม'].map(clean).eq(g)
        if m.any():seen_g.add(g);rows.append([n,g,*metrics(data[m])]);n+=1
    def add_faculty(g,f):
        nonlocal n
        key=(g,f);m=data['__กลุ่ม'].map(clean).eq(g)&data['__คณะ'].map(clean).eq(f)
        if f and key not in seen_f and m.any():seen_f.add(key);rows.append([n,'    '+f,*metrics(data[m])]);n+=1
    def add_program(g,f,p):
        nonlocal n
        key=(g,f,p);m=data['__กลุ่ม'].map(clean).eq(g)&data['__คณะ'].map(clean).eq(f)&data['__สาขา'].map(clean).eq(p)
        if p and key not in seen_p and m.any():seen_p.add(key);rows.append([n,'        '+p,*metrics(data[m])]);n+=1
    for g,faculties in ORDER_TREE:
        add_group(g)
        for f,programs in faculties:
            add_faculty(g,f)
            for p in programs:add_program(g,f,p)
    for _,r in data.sort_values('__ลำดับ').iterrows():
        g=clean(r['__กลุ่ม']);f=clean(r['__คณะ']);p=clean(r['__สาขา'])
        if g not in seen_g:add_group(g)
        if g and f and (g,f) not in seen_f:add_faculty(g,f)
        if g and f and p and (g,f,p) not in seen_p:add_program(g,f,p)
    if rows:rows.append(['','รวมทั้งหมด',*metrics(data)])
    return pd.DataFrame(rows,columns=cols)

def q_summary(data):
    def vals(level):
        x=data[data['__ระดับ']==level];q1=int((x['__Q']=='Q1').sum());q2=int((x['__Q']=='Q2').sum());return q1,q2,q1+q2
    m=vals('ปริญญาโท');d=vals('ปริญญาเอก');t=(m[0]+d[0],m[1]+d[1],m[2]+d[2])
    return pd.DataFrame([['ระดับปริญญาโท',*m],['ระดับปริญญาเอก',*d],['รวม',*t]],columns=['','Scopus Q1','Scopus Q2','รวม'])

def excel_bytes(master,doctor,data):
    summary=q_summary(data);bio=io.BytesIO()
    with pd.ExcelWriter(bio,engine='openpyxl') as w:
        master.to_excel(w,sheet_name='ปริญญาโท',index=False,startrow=6,header=False);doctor.to_excel(w,sheet_name='ปริญญาเอก',index=False,startrow=6,header=False);summary.to_excel(w,sheet_name='สรุป Scopus Q1-Q2',index=False,startrow=2,header=False)
    wb=load_workbook(bio);dark='2F5D1E';light='5A8C3A';white='FFFFFF';peach='FCE4D6';green='E2F0D9';grid='808080';thin=Side(style='thin',color=grid);border=Border(left=thin,right=thin,top=thin,bottom=thin)
    title='ผลงานของนักศึกษาและผู้สำเร็จการศึกษาในระดับ {level} ปีการศึกษา 2568 รอบ 12 เดือน ข้อมูล ระหว่างเดือนกรกฎาคม 2568 ถึง มิถุนายน 2569'
    for ws,level in [(wb['ปริญญาโท'],'ปริญญาโท'),(wb['ปริญญาเอก'],'ปริญญาเอก')]:
        ws.merge_cells('A1:G1');ws['A1']=title.format(level=level)
        for rng in ['A2:A6','B2:B6','C2:D2','C3:C6','D3:D6','E2:E6','F2:G2','F3:F6','G3:G6']:ws.merge_cells(rng)
        labels={'A2':'ลำดับ','B2':'กลุ่ม/คณะ/สาขา','C2':'ระบบ','C3':'ระบบในเวลา\nราชการ','D3':'ระบบนอกเวลา\nราชการ','E2':'จำนวน\nผู้สำเร็จ\nการศึกษา (A)','F2':'1.2.4 จำนวนผลงานระดับบัณฑิตศึกษาที่สามารถตีพิมพ์ใน\nระดับนานาชาติ SCOPUS Q1-Q2','F3':'รวมจำนวนผลงานตีพิมพ์\nในวารสารระดับ\nนานาชาติ Q1-2','G3':'ร้อยละผลงานวิจัยรวมจำนวน\nผลงานตีพิมพ์ในวารสารระดับ\nนานาชาติ Q1-2'}
        for cell,val in labels.items():ws[cell]=val
        for row in range(2,7):
            for col in range(1,8):
                c=ws.cell(row,col);c.fill=PatternFill('solid',fgColor=dark if row<=3 else light);c.font=Font(name='Tahoma',size=12,bold=True,color=white);c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True);c.border=border
        for row in ws.iter_rows(min_row=7,max_row=ws.max_row,min_col=1,max_col=7):
            text=str(row[1].value or '');leading=len(text)-len(text.lstrip());indent=2 if leading>=8 else (1 if leading>=4 else 0);is_group=indent==0 and text.strip() not in ['รวมทั้งหมด',''];fill=peach if is_group or text.strip()=='รวมทั้งหมด' else green
            for c in row:c.border=border;c.font=Font(name='Tahoma',size=12);c.alignment=Alignment(vertical='center',wrap_text=True);c.fill=PatternFill('solid',fgColor=fill)
            row[1].alignment=Alignment(horizontal='left',vertical='center',indent=indent,wrap_text=True)
            for c in row[2:]:c.alignment=Alignment(horizontal='right',vertical='center')
            row[6].number_format='0.00'
        for col,width in {'A':9,'B':52,'C':17,'D':19,'E':18,'F':23,'G':27}.items():ws.column_dimensions[col].width=width
        ws.freeze_panes='A7';ws.sheet_view.showGridLines=False
    ws=wb['สรุป Scopus Q1-Q2'];ws.merge_cells('A1:D1');ws['A1']='1.2.4) จำนวนผลงานระดับบัณฑิตศึกษาที่ตีพิมพ์ในระดับนานาชาติ Scopus Q1, Q2';ws['A1'].font=Font(name='Tahoma',size=14,bold=True);ws['A1'].alignment=Alignment(horizontal='left',vertical='center')
    headers=['','Scopus Q1','Scopus Q2','รวม']
    for c,v in enumerate(headers,1):
        cell=ws.cell(3,c);cell.value=v;cell.font=Font(name='Tahoma',size=12,bold=True);cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True);cell.border=border
    for r in range(4,7):
        for c in range(1,5):
            cell=ws.cell(r,c);cell.border=border;cell.font=Font(name='Tahoma',size=12);cell.alignment=Alignment(horizontal='center' if c>1 else 'left',vertical='center')
    for r,row in enumerate(summary.itertuples(index=False),4):
        for c,v in enumerate(row,1):ws.cell(r,c).value=v
    ws.column_dimensions['A'].width=38;ws.column_dimensions['B'].width=18;ws.column_dimensions['C'].width=18;ws.column_dimensions['D'].width=18
    ws.row_dimensions[1].height=30
    for r in range(3,7):ws.row_dimensions[r].height=30
    ws.freeze_panes='A4';ws.sheet_view.showGridLines=False
    out=io.BytesIO();wb.save(out);out.seek(0);return out

st.title('📊 KPI01 — ตารางสถิติ');st.caption("กลุ่มยึดจาก 'กลุ่มสาขา' | คณะ 1 คณะอยู่ 1 กลุ่ม | สาขา 1 สาขาอยู่ 1 กลุ่ม | ระบบไม่นับรหัสซ้ำ | Q1-Q2 = SCOPUS (Q1)/(Q2)")
uploaded=st.file_uploader('อัปโหลด Excel',type=['xlsx','xls'])
if uploaded:
    try:
        sheet,raw=read_excel(uploaded);data=prepare(raw);master=hierarchy(data[data['__ระดับ']=='ปริญญาโท']);doctor=hierarchy(data[data['__ระดับ']=='ปริญญาเอก']);summary=q_summary(data)
        st.success(f'อ่าน Sheet: {sheet} | ข้อมูล: {len(data):,} แถว')
        st.subheader('ปริญญาโท');st.dataframe(master,use_container_width=True,hide_index=True);st.subheader('ปริญญาเอก');st.dataframe(doctor,use_container_width=True,hide_index=True);st.subheader('สรุปจำนวน Scopus Q1-Q2');st.dataframe(summary,use_container_width=True,hide_index=True)
        st.download_button('📥 ดาวน์โหลด Excel KPI01',excel_bytes(master,doctor,data),'KPI01.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e:st.error(f'ไม่สามารถอ่านไฟล์ได้: {e}')
else:st.info('อัปโหลด Excel เพื่อสร้างตารางสถิติ')
